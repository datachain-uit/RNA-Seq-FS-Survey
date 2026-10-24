"""Dựng RQ0-prime-performance-ceiling.xlsx từ kết quả của run_rq0_prime.py.

Thay artifact_build_20261019/build_rq0_prime_workbook.mjs (bản đó đọc cột ceiling_rule
theo từng phương pháp, đã bỏ khi sửa O14). Sheet đầu tiên là kết luận trần ở mức cohort.

Chạy: python scripts/build_rq0_prime_workbook.py
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

try:  # import khi chạy test (scripts.<x>) hoặc chạy trực tiếp (python scripts/<x>.py)
    from scripts.provenance import DATA_ROOT
except ImportError:
    import sys as _sys
    _sys.path.insert(0, str(Path(__file__).resolve().parent))
    from provenance import DATA_ROOT

ROOT = DATA_ROOT  # thư mục dữ liệu/kết quả (KLTN_DATA_ROOT)
BASE = ROOT / "results" / "phase 1" / "rq0_prime"

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FILL = PatternFill("solid", fgColor="D9EAF7")
YES_FILL = PatternFill("solid", fgColor="FCE4D6")
NO_FILL = PatternFill("solid", fgColor="E2F0D9")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Arial"


def _title(ws, text: str, note: str, width: int) -> None:
    last = get_column_letter(width)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 48


def _table(ws, df: pd.DataFrame, start_row: int, name: str, widths: dict, number_cols: set) -> int:
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(start_row, j, col)
        c.font = Font(name=FONT, size=10, bold=True, color="17365D")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=1):
            c = ws.cell(i, j, value.item() if hasattr(value, "item") else value)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if df.columns[j - 1] in number_cols:
                c.number_format = "0.000"
    end = start_row + len(df)
    ref = f"A{start_row}:{get_column_letter(len(df.columns))}{end}"
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    return end


def build() -> Path:
    config = json.loads((BASE / "config.json").read_text(encoding="utf-8"))
    summary = pd.read_csv(BASE / "summary.csv", encoding="utf-8-sig", dtype={"train_per_class": str})
    ceiling = pd.read_csv(BASE / "ceiling_by_cohort.csv", encoding="utf-8-sig", dtype={"train_per_class": str})

    g = next(d for d in config["input_datasets"] if d["cohort"] == "gse30219")
    if g["n_samples"] != 146 or g.get("class_counts", {}) != {"LUAD": 85, "LUSC": 61}:
        raise ValueError("Cần GSE30219 đã sửa R5 (146 mẫu: 85 LUAD / 61 LUSC).")
    if "ceiling_rule" in summary.columns:
        raise ValueError("summary.csv còn cột ceiling_rule cũ; chạy lại scripts/run_rq0_prime.py.")

    rule = config["ceiling_rule"]
    wb = Workbook()

    # Sheet 1: kết luận trần mức cohort
    ws = wb.active
    ws.title = "Kết luận trần"
    view = pd.DataFrame({
        "Cohort": ceiling["cohort"],
        "AUROC gen đơn tốt nhất (n_T = 5)": ceiling["auroc_best_single_gene"],
        "AUROC IHC-2 (n_T = 5)": ceiling["auroc_IHC_2"],
        "AUROC nhỏ nhất": ceiling["min_auroc"],
        "Phương pháp thấp nhất": ceiling["min_method"],
        "Ngưỡng": ceiling["threshold"],
        "Chạm trần": ceiling["ceiling_hit"].map({True: "Có", False: "Không"}),
    })
    hits = ceiling.loc[ceiling["ceiling_hit"], "cohort"].tolist()
    verdict = ("Có cohort chạm trần: " + ", ".join(hits)) if hits else \
        "Không cohort nào chạm trần → giữ nhiệm vụ chính LUAD vs LUSC, không mở nhiệm vụ dự phòng."
    _title(ws, "RQ0′ — Kết luận trần hiệu năng theo cohort",
           f"Quy tắc (sửa O14, Review 20/10/2026): ở n_T = {rule['train_per_class']} mẫu mỗi lớp, cohort chạm trần nếu "
           f"mọi phương pháp trong {{{', '.join(rule['methods'])}}} đạt AUROC trung bình ≥ {rule['auroc_mean_threshold']}. "
           "B0, IHC-4 và random-k không thuộc tập xét. CV 5×5; GSE30219 dùng bản đã sửa R5 (85/61). "
           f"Tạo lúc {config['generated_at']}.", len(view.columns))
    end = _table(ws, view, 4, "RQ0Ceiling",
                 {"A": 14, "B": 20, "C": 18, "D": 16, "E": 22, "F": 10, "G": 12},
                 {"AUROC gen đơn tốt nhất (n_T = 5)", "AUROC IHC-2 (n_T = 5)", "AUROC nhỏ nhất", "Ngưỡng"})
    rng = f"G5:G{end}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$G5="Có"'], fill=YES_FILL, font=Font(color="C00000", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$G5="Không"'], fill=NO_FILL, font=Font(color="375623", bold=True)))
    ws.cell(end + 2, 1, "Kết luận:").font = Font(name=FONT, size=11, bold=True)
    ws.merge_cells(start_row=end + 2, start_column=2, end_row=end + 2, end_column=len(view.columns))
    ws.cell(end + 2, 2, verdict).font = Font(name=FONT, size=11, bold=True, color="17365D")
    ws.cell(end + 3, 1, "Ghi chú:").font = Font(name=FONT, size=10, italic=True)
    ws.merge_cells(start_row=end + 3, start_column=2, end_row=end + 3, end_column=len(view.columns))
    ws.cell(end + 3, 2, "25 fold không độc lập (tập train chồng lấn); so sánh phương pháp phải dùng phương sai hiệu "
                        "chỉnh Nadeau–Bengio. Nguồn: ceiling_by_cohort.csv.").font = Font(name=FONT, size=10, italic=True)
    ws.cell(end + 3, 2).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[end + 3].height = 30

    # Sheet 2: toàn bộ summary
    ws2 = wb.create_sheet("RQ0-prime summary")
    _title(ws2, "RQ0′ — Bảng tổng hợp (trung bình ± SD qua 25 fold)",
           f"Repeated stratified 5×5 CV. GSE30219 dùng {g['data_directory']} ({g['n_samples']} mẫu: "
           f"{g['class_counts']['LUAD']} LUAD / {g['class_counts']['LUSC']} LUSC). Tạo lúc {config['generated_at']}.",
           len(summary.columns))
    _table(ws2, summary, 4, "RQ0PrimeSummary",
           {"A": 14, "B": 15, "C": 20, "D": 9, "E": 12, "F": 12, "G": 12, "H": 12},
           {"auroc_mean", "auroc_sd", "bacc_mean", "bacc_sd"})

    for sheet in (ws, ws2):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    out = BASE / "RQ0-prime-performance-ceiling.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    print(f"Wrote {build()}")
